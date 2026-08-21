#!/usr/bin/env python3
"""
run_openmp_bench.py
====================

Runs every OpenMP benchmark binary from GPU4S Bench across a set of thread
counts, repeats each run 3x (configurable), averages the timing, and writes
the averages straight into a copy of the tracking spreadsheet
(input_file_to_be_filled.xlsx-style layout).

Supports both local execution and execution on an Android device via ADB.

WHAT IT MEASURES
----------------
Every binary is run with the `-c` flag, which makes the GPU4S OpenMP
binaries print a line of three semicolon-separated numbers, e.g.:

    0.0000000000;4580.3461914062;0.0000000000;

This script parses that line and takes the **middle value** (the actual compute
time) as the measurement for each run.

USAGE (Android via ADB)
-----------------------
    python3 run_openmp_bench.py \
        --use-adb \
        --bin-dir /data/local/tmp \
        --input input_file_to_be_filled.xlsx \
        --output GPU4S_bench_result_filled.xlsx \
        --threads 1 2 4 8 16 \
        --repeats 3 \
        --pause 15
"""

import argparse
import csv
import re
import statistics
import subprocess
import sys
import time
from pathlib import Path

import openpyxl

# ======================================================================
# Benchmark -> binary name mapping
# ======================================================================
BENCH_PREFIX = {
    "CIFAR-10": "cifar_10",
    "CIFAR-10-Multiple": "cifar_10_multiple",
    "2D Convolution": "convolution_2D",
    "2D Correlation": "correlation_2D",
    "FFT": "FFT",
    "Window FFT": "FFT_window",
    "FIR filter": "FIR_filter",
    "LRN": "LRN",
    "Matrix Multiplication": "matrix_mult",
    "Max pooling": "max_pooling",
    "ReLU": "relu",
    "Softmax": "softmax",
    "Wavelet transform": "wavelet_transform",
}

VERSION_SUFFIX = {
    "Naïve": "",
    "Optimized": "_opt",
    "Lib": "_lib",
}

TIMING_LINE_RE = re.compile(
    r"([-+]?\d+(?:\.\d+)?)\s*;\s*([-+]?\d+(?:\.\d+)?)\s*;\s*([-+]?\d+(?:\.\d+)?)\s*;?"
)


def sanitize_config(raw_config):
    """
    Collapses spaces used as thousands separators in config strings
    and returns argument list.
    """
    if not raw_config or str(raw_config).strip().lower() == "no option":
        return []
    collapsed = re.sub(r"(?<=\d)\s(?=\d)", "", str(raw_config))
    return collapsed.split()


def build_thread_column_map(ws, header_row=2):
    """
    Reads header row and returns {thread_count: column_index} for 'OpenMP-<N>'.
    """
    mapping = {}
    for col in range(1, ws.max_column + 1):
        val = ws.cell(row=header_row, column=col).value
        if val and str(val).startswith("OpenMP-"):
            try:
                n = int(str(val).split("-")[1])
                mapping[n] = col
            except (ValueError, IndexError):
                continue
    return mapping


def resolve_merged(ws):
    """
    Fills merged cell values across their ranges.
    """
    merged_map = {}
    for mrange in ws.merged_cells.ranges:
        top_val = ws.cell(row=mrange.min_row, column=mrange.min_col).value
        for r in range(mrange.min_row, mrange.max_row + 1):
            for c in range(mrange.min_col, mrange.max_col + 1):
                merged_map[(r, c)] = top_val
    return merged_map


def get_cell(ws, merged_map, row, col):
    return merged_map.get((row, col), ws.cell(row=row, column=col).value)


def run_once(binary_str, args, thread_count, env_base, use_adb=False, adb_device=None):
    """
    Runs the binary locally or over ADB with OMP_NUM_THREADS=thread_count and -c flag.
    Returns (stdout_text, returncode).
    """
    full_args = list(args)
    if "-c" not in full_args:
        full_args.append("-c")

    if use_adb:
        adb_prefix = ["adb"]
        if adb_device:
            adb_prefix.extend(["-s", adb_device])
        
        args_str = " ".join(full_args)
        remote_cmd = f"OMP_NUM_THREADS={thread_count} {binary_str} {args_str}"
        cmd = adb_prefix + ["shell", remote_cmd]
        
        try:
            result = subprocess.run(
                cmd, capture_output=True, text=True, timeout=1800
            )
            return result.stdout + result.stderr, result.returncode
        except subprocess.TimeoutExpired:
            return "TIMEOUT", -1
    else:
        env = dict(env_base)
        env["OMP_NUM_THREADS"] = str(thread_count)
        cmd = [binary_str] + full_args

        try:
            result = subprocess.run(
                cmd, env=env, capture_output=True, text=True, timeout=1800
            )
            return result.stdout + result.stderr, result.returncode
        except subprocess.TimeoutExpired:
            return "TIMEOUT", -1


def parse_internal_timing(stdout_text):
    """
    Parses <init>;<compute>;<transfer>; output and returns the middle value.
    """
    matches = TIMING_LINE_RE.findall(stdout_text)
    if not matches:
        return None
    _, middle, _ = matches[-1]
    try:
        return float(middle)
    except ValueError:
        return None


def check_binary_exists(bin_path_str, use_adb=False, adb_device=None):
    """
    Checks binary existence locally or on remote Android device.
    """
    if use_adb:
        adb_prefix = ["adb"]
        if adb_device:
            adb_prefix.extend(["-s", adb_device])
        cmd = adb_prefix + ["shell", f"[ -f {bin_path_str} ] && echo EXISTS"]
        res = subprocess.run(cmd, capture_output=True, text=True)
        return "EXISTS" in res.stdout
    else:
        return Path(bin_path_str).is_file()


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--bin-dir", required=True, help="Directory containing binaries (e.g. /data/local/tmp or ./bin)")
    ap.add_argument("--input", required=True, help="Input xlsx template")
    ap.add_argument("--output", required=True, help="Output xlsx")
    ap.add_argument("--threads", type=int, nargs="+", default=[1, 2, 4, 8, 16],
                     help="Thread counts to test")
    ap.add_argument("--repeats", type=int, default=3, help="Number of repeats per test")
    ap.add_argument("--pause", type=float, default=15.0, help="Cooldown pause in seconds")
    ap.add_argument("--sheet", default="Data", help="Worksheet name")
    ap.add_argument("--use-adb", action="store_true", help="Run binaries on Android via ADB shell")
    ap.add_argument("--adb-device", default=None, help="Specific ADB device serial (optional)")
    ap.add_argument("--dry-run", action="store_true", help="Print planned commands, run nothing")
    ap.add_argument("--force", action="store_true", help="Overwrite cells that already have a value")
    args = ap.parse_args()

    input_path = Path(args.input)
    output_path = Path(args.output)
    log_path = output_path.with_suffix(".raw_runs.csv")

    if not args.use_adb and not Path(args.bin_dir).is_dir():
        sys.exit(f"ERROR: Local bin-dir not found: {args.bin_dir}")
    if not input_path.is_file():
        sys.exit(f"ERROR: Input xlsx not found: {input_path}")

    wb = openpyxl.load_workbook(input_path)
    ws = wb[args.sheet]

    thread_cols = build_thread_column_map(ws)
    missing = [t for t in args.threads if t not in thread_cols]
    if missing:
        sys.exit(f"ERROR: sheet missing OpenMP-{missing} column(s). Found: {sorted(thread_cols)}")

    merged_map = resolve_merged(ws)

    tasks = []
    for row in range(3, ws.max_row + 1):
        bench = get_cell(ws, merged_map, row, 2)
        version = ws.cell(row=row, column=3).value
        if not bench or not version:
            continue
        raw_config = get_cell(ws, merged_map, row, 10)
        tasks.append({
            "row": row,
            "bench": bench,
            "version": version,
            "args": sanitize_config(raw_config),
        })

    log_exists = log_path.is_file()
    log_file = open(log_path, "a", newline="")
    log_writer = csv.writer(log_file)
    if not log_exists:
        log_writer.writerow(["benchmark", "version", "binary", "threads", "run_index",
                              "timing_ms", "returncode", "status"])

    env_base = dict(__import__("os").environ)

    total_planned = 0
    plan = []
    bin_dir_clean = args.bin_dir.rstrip("/")
    
    for task in tasks:
        prefix = BENCH_PREFIX.get(task["bench"])
        suffix = VERSION_SUFFIX.get(task["version"])
        if prefix is None or suffix is None:
            print(f"WARNING: no binary mapping for '{task['bench']}' / '{task['version']}' — skipping row {task['row']}")
            continue
        
        binary_name = f"{prefix}_openmp{suffix}"
        binary_path_str = f"{bin_dir_clean}/{binary_name}"

        if not check_binary_exists(binary_path_str, use_adb=args.use_adb, adb_device=args.adb_device):
            print(f"WARNING: binary not found, skipping: {binary_path_str}")
            continue

        for t in args.threads:
            col = thread_cols[t]
            existing = ws.cell(row=task["row"], column=col).value
            if existing is not None and not args.force:
                continue
            plan.append({**task, "binary_str": binary_path_str, "binary_name": binary_name, "threads": t, "col": col})
            total_planned += 1

    print(f"Mode: {'ADB (Android Device)' if args.use_adb else 'Local Execution'}")
    print(f"Planned runs: {total_planned} cells x {args.repeats} repeats "
          f"= {total_planned * args.repeats} benchmark executions")
    print(f"Estimated minimum cooldown time: "
          f"{total_planned * args.repeats * args.pause / 60:.1f} minutes")

    if args.dry_run:
        for p in plan:
            print(f"[DRY RUN] {p['binary_name']} threads={p['threads']} args={p['args']}")
        return

    for i, p in enumerate(plan, start=1):
        print(f"\n[{i}/{total_planned}] {p['bench']} / {p['version']} "
              f"({p['binary_name']}) threads={p['threads']} args={' '.join(p['args']) or '(none)'}")

        samples = []
        for rep in range(1, args.repeats + 1):
            output_text, rc = run_once(
                p["binary_str"], p["args"], p["threads"], env_base, 
                use_adb=args.use_adb, adb_device=args.adb_device
            )

            timing = parse_internal_timing(output_text)

            if timing is None:
                status = "PARSE_FAILED"
                print(f"    repeat {rep}/{args.repeats}: FAILED to parse timing line "
                      f"(returncode={rc}) — output: {output_text[:150]!r}")
                log_writer.writerow([p["bench"], p["version"], p["binary_name"],
                                      p["threads"], rep, "", rc, status])
            else:
                status = "OK" if rc == 0 else "OK_NONZERO_RC"
                print(f"    repeat {rep}/{args.repeats}: {timing:.4f} ms (returncode={rc})")
                samples.append(timing)
                log_writer.writerow([p["bench"], p["version"], p["binary_name"],
                                      p["threads"], rep, f"{timing:.4f}", rc, status])
            log_file.flush()

            print(f"    cooling down {args.pause:.0f}s...")
            time.sleep(args.pause)

        if not samples:
            print(f"  -> ALL {args.repeats} repeats failed to parse — leaving cell empty.")
            continue

        avg_ms = statistics.mean(samples)
        ws.cell(row=p["row"], column=p["col"], value=round(avg_ms, 4))
        print(f"  -> average: {avg_ms:.4f} ms  (from {samples})")

        wb.save(output_path)

    log_file.close()
    print(f"\nDone. Filled workbook: {output_path}")
    print(f"Raw per-run log:       {log_path}")


if __name__ == "__main__":
    main()