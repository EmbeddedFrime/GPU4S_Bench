#!/usr/bin/env python3
"""
run_allsuite_bench_phone.py
============================

Runs the FULL GPU4S Bench suite (CPU, OpenMP, OpenCL) on an Android phone
(tested against the "Phone (Snapdragon 810 + Adreno 430)" tracking sheet)
and writes the H2D / Kernel / D2H timings straight into a copy of the
tracking spreadsheet.

Unlike run_openmp_bench.py (which only swept OpenMP thread counts), this
sheet has one row per Framework/Version combo and THREE timing columns:

    H2D Transfer (ms) | Kernel Execution (ms) | D2H Transfer (ms)

("Total (ms)" is a live =SUM() formula in the sheet and is left alone.)

Expected sheet layout (header row 2, data from row 3):
    A: N°            B: Benchmark        C: Framework   D: Version
    E: H2D Transfer  F: Kernel Execution G: D2H Transfer
    H: Total (formula, not written)      I: Notes        J: Config args

Every binary is still expected to be run with `-c` and to print a line of
three semicolon-separated numbers:

    <h2d_ms>;<kernel_ms>;<d2h_ms>;

USAGE (Android via ADB, the default target for this sheet)
------------------------------------------------------------
    python3 run_allsuite_bench_phone.py \
        --use-adb \
        --bin-dir /data/local/tmp \
        --input input_file_allsuite_phone.xlsx \
        --output GPU4S_allsuite_phone_result.xlsx \
        --omp-threads 8 \
        --repeats 3 \
        --pause 15

USAGE (local execution, e.g. testing on a laptop first)
------------------------------------------------------------
    python3 run_allsuite_bench_phone.py \
        --bin-dir ./bin \
        --input input_file_allsuite_phone.xlsx \
        --output GPU4S_allsuite_local_result.xlsx \
        --dry-run

NOTE ON BINARY NAMING
----------------------
Binary names are assembled as:  <bench_prefix><framework_suffix><version_suffix>
"+ UMA" versions are NOT separate binaries — they reuse the non-UMA binary
for that Version and are launched with an extra `-u` flag instead.
The BENCH_PREFIX / FRAMEWORK_SUFFIX / VERSION_INFO tables below are best
guesses that follow the same convention as the existing OpenMP-only script.
Run with --dry-run first (it prints every resolved binary name + args, and
checks whether the binary exists in --bin-dir) and adjust the tables if any
of your actual binaries are named differently.
"""

import argparse
import csv
import re
import statistics
import subprocess
import sys
import time
from pathlib import Path

import openpyxl

# ======================================================================
# Benchmark -> binary prefix mapping (adjust to match your actual binaries)
# ======================================================================
BENCH_PREFIX = {
    "CIFAR-10": "cifar_10",
    "CIFAR-10-Multiple": "cifar_10_multiple",
    "2D Convolution": "convolution_2D",
    "2D Correlation": "correlation_2D",
    "2D FFT": "FFT_2D",
    "FFT": "FFT",
    "Window FFT": "FFT_window",
    "FIR filter": "FIR_filter",
    "LRN": "LRN",
    "Matrix Multiplication": "matrix_mult",
    "Matrix Multiplication FP16": "matrix_mult_fp16",
    "Matrix Multiplication tensor": "matrix_mult_tensor",
    "Max pooling": "max_pooling",
    "Memory bandwidth": "memory_bandwidth",
    "ReLU": "relu",
    "Softmax": "softmax",
    "Wavelet transform": "wavelet_transform",
}

# Framework -> binary suffix
FRAMEWORK_SUFFIX = {
    "CPU": "_cpu",
    "OpenMP": "_openmp",
    "OpenCL": "_opencl",
}

# Version -> (binary suffix, extra CLI args)
# NOTE: "+ UMA" variants are NOT separate binaries — they're the same binary
# as their non-UMA counterpart, launched with an extra -u flag.
VERSION_INFO = {
    "Naïve": {"suffix": "", "extra_args": []},
    "Naïve + UMA": {"suffix": "", "extra_args": ["-u"]},
    "Optimized": {"suffix": "_opt", "extra_args": []},
    "Optimized + UMA": {"suffix": "_opt", "extra_args": ["-u"]},
    "Lib": {"suffix": "_lib", "extra_args": []},
    "Lib + UMA": {"suffix": "_lib", "extra_args": ["-u"]},
}

# Frameworks that should be launched with OMP_NUM_THREADS set
THREADED_FRAMEWORKS = {"OpenMP"}

TIMING_LINE_RE = re.compile(
    r"([-+]?\d+(?:\.\d+)?)\s*;\s*([-+]?\d+(?:\.\d+)?)\s*;\s*([-+]?\d+(?:\.\d+)?)\s*;?"
)

# Column indices (1-based) for the fixed sheet layout described above
COL_N = 1
COL_BENCH = 2
COL_FRAMEWORK = 3
COL_VERSION = 4
COL_H2D = 5
COL_KERNEL = 6
COL_D2H = 7
COL_TOTAL = 8   # formula - never written
COL_NOTES = 9
COL_CONFIG = 10


def sanitize_config(raw_config):
    """Collapses spaces used as thousands separators in config strings."""
    if not raw_config or str(raw_config).strip().lower() == "no option":
        return []
    collapsed = re.sub(r"(?<=\d)\s(?=\d)", "", str(raw_config))
    return collapsed.split()


def resolve_merged(ws):
    """Fills merged cell values across their ranges."""
    merged_map = {}
    for mrange in ws.merged_cells.ranges:
        top_val = ws.cell(row=mrange.min_row, column=mrange.min_col).value
        for r in range(mrange.min_row, mrange.max_row + 1):
            for c in range(mrange.min_col, mrange.max_col + 1):
                merged_map[(r, c)] = top_val
    return merged_map


def get_cell(ws, merged_map, row, col):
    return merged_map.get((row, col), ws.cell(row=row, column=col).value)


def check_binary_exists(bin_path_str, use_adb=False, adb_device=None):
    if use_adb:
        adb_prefix = ["adb"]
        if adb_device:
            adb_prefix.extend(["-s", adb_device])
        cmd = adb_prefix + ["shell", f"[ -f {bin_path_str} ] && echo EXISTS"]
        res = subprocess.run(cmd, capture_output=True, text=True)
        return "EXISTS" in res.stdout
    else:
        return Path(bin_path_str).is_file()


def run_once(binary_str, args, env_extra, use_adb=False, adb_device=None):
    """
    Runs the binary locally or over ADB with the given env vars and -c flag.
    Returns (stdout_text, returncode).
    """
    full_args = list(args)
    if "-c" not in full_args:
        full_args.append("-c")

    if use_adb:
        adb_prefix = ["adb"]
        if adb_device:
            adb_prefix.extend(["-s", adb_device])

        env_str = " ".join(f"{k}={v}" for k, v in env_extra.items())
        args_str = " ".join(full_args)
        remote_cmd = f"{env_str} {binary_str} {args_str}".strip()
        cmd = adb_prefix + ["shell", remote_cmd]

        try:
            result = subprocess.run(cmd, capture_output=True, text=True, timeout=1800)
            return result.stdout + result.stderr, result.returncode
        except subprocess.TimeoutExpired:
            return "TIMEOUT", -1
    else:
        import os
        env = dict(os.environ)
        env.update(env_extra)
        cmd = [binary_str] + full_args
        try:
            result = subprocess.run(cmd, env=env, capture_output=True, text=True, timeout=1800)
            return result.stdout + result.stderr, result.returncode
        except subprocess.TimeoutExpired:
            return "TIMEOUT", -1


def parse_internal_timing(stdout_text):
    """Parses <h2d>;<kernel>;<d2h>; and returns (h2d, kernel, d2h) floats."""
    matches = TIMING_LINE_RE.findall(stdout_text)
    if not matches:
        return None
    h2d, kernel, d2h = matches[-1]
    try:
        return float(h2d), float(kernel), float(d2h)
    except ValueError:
        return None


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--bin-dir", required=True, help="Directory containing binaries (e.g. /data/local/tmp or ./bin)")
    ap.add_argument("--input", required=True, help="Input xlsx template")
    ap.add_argument("--output", required=True, help="Output xlsx")
    ap.add_argument("--omp-threads", type=int, default=8, help="OMP_NUM_THREADS for OpenMP rows")
    ap.add_argument("--repeats", type=int, default=3, help="Number of repeats per row")
    ap.add_argument("--pause", type=float, default=15.0, help="Cooldown pause in seconds between runs")
    ap.add_argument("--sheet", default="Data", help="Worksheet name")
    ap.add_argument("--use-adb", action="store_true", help="Run binaries on the phone via ADB shell (recommended)")
    ap.add_argument("--adb-device", default=None, help="Specific ADB device serial (optional)")
    ap.add_argument("--dry-run", action="store_true", help="Print planned commands / resolved binaries, run nothing")
    ap.add_argument("--force", action="store_true", help="Overwrite cells that already have a value")
    ap.add_argument("--only-framework", choices=["CPU", "OpenMP", "OpenCL"], default=None,
                     help="Restrict this run to a single framework (useful for re-running just one)")
    args = ap.parse_args()

    input_path = Path(args.input)
    output_path = Path(args.output)
    log_path = output_path.with_suffix(".raw_runs.csv")

    if not args.use_adb and not Path(args.bin_dir).is_dir():
        sys.exit(f"ERROR: Local bin-dir not found: {args.bin_dir}")
    if not input_path.is_file():
        sys.exit(f"ERROR: Input xlsx not found: {input_path}")

    wb = openpyxl.load_workbook(input_path)
    ws = wb[args.sheet]
    merged_map = resolve_merged(ws)

    bin_dir_clean = args.bin_dir.rstrip("/")

    plan = []
    for row in range(3, ws.max_row + 1):
        bench = get_cell(ws, merged_map, row, COL_BENCH)
        framework = get_cell(ws, merged_map, row, COL_FRAMEWORK)
        version = ws.cell(row=row, column=COL_VERSION).value
        if not bench or not framework or not version:
            continue
        if args.only_framework and framework != args.only_framework:
            continue

        prefix = BENCH_PREFIX.get(bench)
        fw_suffix = FRAMEWORK_SUFFIX.get(framework)
        ver_info = VERSION_INFO.get(version)
        if prefix is None or fw_suffix is None or ver_info is None:
            print(f"WARNING: no binary mapping for '{bench}' / '{framework}' / '{version}' "
                  f"— skipping row {row}")
            continue

        binary_name = f"{prefix}{fw_suffix}{ver_info['suffix']}"
        binary_path_str = f"{bin_dir_clean}/{binary_name}"

        # Skip only if ALL three timing cells are already filled (unless --force)
        existing = (
            ws.cell(row=row, column=COL_H2D).value,
            ws.cell(row=row, column=COL_KERNEL).value,
            ws.cell(row=row, column=COL_D2H).value,
        )
        if all(v is not None for v in existing) and not args.force:
            continue

        raw_config = get_cell(ws, merged_map, row, COL_CONFIG)
        env_extra = {}
        if framework in THREADED_FRAMEWORKS:
            env_extra["OMP_NUM_THREADS"] = str(args.omp_threads)

        row_args = sanitize_config(raw_config) + ver_info["extra_args"]

        plan.append({
            "row": row,
            "bench": bench,
            "framework": framework,
            "version": version,
            "args": row_args,
            "binary_name": binary_name,
            "binary_str": binary_path_str,
            "env_extra": env_extra,
        })

    # Check existence up front so a missing binary doesn't burn a cooldown cycle
    missing_binaries = []
    for p in plan:
        if not check_binary_exists(p["binary_str"], use_adb=args.use_adb, adb_device=args.adb_device):
            missing_binaries.append(p["binary_str"])
    if missing_binaries:
        print(f"WARNING: {len(missing_binaries)} binaries not found on target, those rows will be skipped:")
        for m in sorted(set(missing_binaries)):
            print(f"    {m}")
    plan = [p for p in plan
            if check_binary_exists(p["binary_str"], use_adb=args.use_adb, adb_device=args.adb_device)]

    print(f"Mode: {'ADB (Android Device)' if args.use_adb else 'Local Execution'}")
    print(f"Planned rows: {len(plan)}  x  {args.repeats} repeats "
          f"= {len(plan) * args.repeats} benchmark executions")
    print(f"Estimated minimum cooldown time: "
          f"{len(plan) * args.repeats * args.pause / 60:.1f} minutes")

    if args.dry_run:
        for p in plan:
            env_str = " ".join(f"{k}={v}" for k, v in p["env_extra"].items())
            print(f"[DRY RUN] {env_str} {p['binary_name']} args={p['args'] or '(none)'} "
                  f"[{p['framework']} / {p['version']}]")
        return

    log_exists = log_path.is_file()
    log_file = open(log_path, "a", newline="")
    log_writer = csv.writer(log_file)
    if not log_exists:
        log_writer.writerow(["benchmark", "framework", "version", "binary", "run_index",
                              "h2d_ms", "kernel_ms", "d2h_ms", "returncode", "status"])

    for i, p in enumerate(plan, start=1):
        print(f"\n[{i}/{len(plan)}] {p['bench']} / {p['framework']} / {p['version']} "
              f"({p['binary_name']}) args={' '.join(p['args']) or '(none)'}")

        h2d_samples, kernel_samples, d2h_samples = [], [], []
        for rep in range(1, args.repeats + 1):
            output_text, rc = run_once(
                p["binary_str"], p["args"], p["env_extra"],
                use_adb=args.use_adb, adb_device=args.adb_device
            )

            timing = parse_internal_timing(output_text)
            if timing is None:
                status = "PARSE_FAILED"
                print(f"    repeat {rep}/{args.repeats}: FAILED to parse timing line "
                      f"(returncode={rc}) — output: {output_text[:150]!r}")
                log_writer.writerow([p["bench"], p["framework"], p["version"], p["binary_name"],
                                      rep, "", "", "", rc, status])
            else:
                h2d, kernel, d2h = timing
                status = "OK" if rc == 0 else "OK_NONZERO_RC"
                print(f"    repeat {rep}/{args.repeats}: h2d={h2d:.4f} kernel={kernel:.4f} "
                      f"d2h={d2h:.4f} ms (returncode={rc})")
                h2d_samples.append(h2d)
                kernel_samples.append(kernel)
                d2h_samples.append(d2h)
                log_writer.writerow([p["bench"], p["framework"], p["version"], p["binary_name"],
                                      rep, f"{h2d:.4f}", f"{kernel:.4f}", f"{d2h:.4f}", rc, status])
            log_file.flush()

            print(f"    cooling down {args.pause:.0f}s...")
            time.sleep(args.pause)

        if not kernel_samples:
            print(f"  -> ALL {args.repeats} repeats failed to parse — leaving row empty.")
            continue

        avg_h2d = statistics.mean(h2d_samples)
        avg_kernel = statistics.mean(kernel_samples)
        avg_d2h = statistics.mean(d2h_samples)
        ws.cell(row=p["row"], column=COL_H2D, value=round(avg_h2d, 4))
        ws.cell(row=p["row"], column=COL_KERNEL, value=round(avg_kernel, 4))
        ws.cell(row=p["row"], column=COL_D2H, value=round(avg_d2h, 4))
        print(f"  -> averages: h2d={avg_h2d:.4f}  kernel={avg_kernel:.4f}  d2h={avg_d2h:.4f} ms "
              f"(from {len(kernel_samples)} valid repeats)")

        wb.save(output_path)

    log_file.close()
    print(f"\nDone. Filled workbook: {output_path}")
    print(f"Raw per-run log:       {log_path}")


if __name__ == "__main__":
    main()